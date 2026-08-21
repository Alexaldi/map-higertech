from pathlib import Path
from docx import Document

path = Path(r'C:\Users\Aspire 7\Downloads\Manual_Book_User_SIH3_Maluku_Final.docx')
doc = Document(path)

print('CORE', doc.core_properties.title, '|', doc.core_properties.author, '|', doc.core_properties.subject)
print('PARAGRAPHS', len(doc.paragraphs), 'TABLES', len(doc.tables), 'SHAPES', len(doc.inline_shapes), 'SECTIONS', len(doc.sections))
for si, section in enumerate(doc.sections, 1):
    print(f'SECTION {si} page={section.page_width.inches:.2f}x{section.page_height.inches:.2f} margins={section.left_margin.inches:.2f},{section.right_margin.inches:.2f},{section.top_margin.inches:.2f},{section.bottom_margin.inches:.2f} header={section.header_distance.inches:.2f} footer={section.footer_distance.inches:.2f} diff_first={section.different_first_page_header_footer}')
    for part_name, part in [('HEADER', section.header), ('FIRST_HEADER', section.first_page_header), ('FOOTER', section.footer), ('FIRST_FOOTER', section.first_page_footer)]:
        texts = [p.text for p in part.paragraphs if p.text.strip()]
        tables = len(part.tables)
        print(f'  {part_name} paragraphs={texts!r} tables={tables}')

for i, paragraph in enumerate(doc.paragraphs, 1):
    text = paragraph.text.replace('\n', ' | ')
    if text.strip() or paragraph.style.name != 'Normal':
        fmt = paragraph.paragraph_format
        print(f'P{i:03d} style={paragraph.style.name!r} align={paragraph.alignment} before={fmt.space_before} after={fmt.space_after} line={fmt.line_spacing} text={text!r}')
        for j, run in enumerate(paragraph.runs, 1):
            if run.text.strip():
                print(f'  R{j} text={run.text!r} font={run.font.name!r} size={run.font.size.pt if run.font.size else None} bold={run.bold} italic={run.italic} color={run.font.color.rgb if run.font.color and run.font.color.type else None}')

for ti, table in enumerate(doc.tables, 1):
    print(f'TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} style={table.style.name if table.style else None}')
    for ri, row in enumerate(table.rows[:12], 1):
        print(f'  ROW {ri}: {[cell.text.replace(chr(10), " | ") for cell in row.cells]!r}')

print('STYLES')
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    pf = style.paragraph_format
    font = style.font
    print(f'STYLE {name!r} font={font.name!r} size={font.size.pt if font.size else None} bold={font.bold} color={font.color.rgb if font.color and font.color.type else None} before={pf.space_before} after={pf.space_after} line={pf.line_spacing}')
