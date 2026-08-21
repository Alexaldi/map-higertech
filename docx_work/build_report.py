from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Aspire 7\Documents\ChatGPT\map-higertech")
WORK = ROOT / "docx_work"
REF = Path(r"C:\Users\Aspire 7\Downloads\Manual_Book_User_SIH3_Maluku_Final.docx")
OUT = ROOT / "Laporan_Aplikasi_Higertech_Live_Monitoring.docx"


def set_run(run, size=None, bold=None, italic=None, color=None, font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def style_paragraph(p, before=None, after=None, line=None, keep=False):
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if keep:
        pPr = p._p.get_or_add_pPr()
        keep_next = OxmlElement("w:keepNext")
        pPr.append(keep_next)


def add_text(doc, text, style="Normal", align=None, before=None, after=None, line=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run(r)
    style_paragraph(p, before=before, after=after, line=line)
    return p


def add_rich(doc, parts, style="Normal", align=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    for text, opts in parts:
        r = p.add_run(text)
        set_run(r, **opts)
    style_paragraph(p, before=before, after=after)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run(r, bold=True)
    style_paragraph(p, keep=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Quote")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=8, italic=True, color=(90, 90, 90))
    return p


def add_image(doc, path, caption, width=6.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    # Add a useful description for screen readers.
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", caption)
        drawing[0].set("name", caption)
    add_caption(doc, caption)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=(35, 52, 78), size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    set_run(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width_inches):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_inches * 1440)))
    tcW.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def add_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(width * 1440)))
        grid.append(gc)
    head = table.rows[0]
    set_repeat_table_header(head)
    for i, text in enumerate(headers):
        set_cell_width(head.cells[i], widths[i])
        set_cell_margins(head.cells[i])
        set_cell_text(head.cells[i], text, bold=True, color=(255, 255, 255), size=font_size)
        shade(head.cells[i], "234E9B")
    for row_data in rows:
        row = table.add_row()
        for i, text in enumerate(row_data):
            set_cell_width(row.cells[i], widths[i])
            set_cell_margins(row.cells[i])
            set_cell_text(row.cells[i], text, size=font_size)
            shade(row.cells[i], "F7F9FC" if len(table.rows) % 2 == 0 else "FFFFFF")
    doc.add_paragraph()
    return table


def add_label_para(doc, label, text):
    return add_rich(doc, [(label, {"bold": True}), (text, {})])


def clear_body(doc):
    body = doc._element.body
    sect = body.sectPr
    for child in list(body):
        if child is not sect:
            body.remove(child)


def set_core_properties(doc):
    props = doc.core_properties
    props.title = "Laporan Aplikasi Higertech Live Monitoring Map"
    props.subject = "Laporan pengembangan prototype WebGIS monitoring telemetri"
    props.author = "Tim Pengembang Higertech Live Monitoring"
    props.keywords = "Laravel, WebGIS, Leaflet, telemetry, monitoring, Higertech"
    props.comments = "Dibuat dari template Manual Book User SIH3 Maluku dengan struktur laporan aplikasi."


def set_alt_text_for_drawings(doc):
    """Keep template images and newly inserted screenshots accessible."""
    parts = [doc.part]
    parts.extend(
        rel.target_part
        for rel in doc.part.rels.values()
        if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer")
    )
    for part in parts:
        root = getattr(part, "_element", None)
        if root is None:
            continue
        for index, doc_pr in enumerate(root.xpath(".//wp:docPr"), start=1):
            descr = doc_pr.get("descr") or doc_pr.get("title")
            if not descr:
                descr = "Higertech Live Monitoring visual asset %d" % index
            doc_pr.set("descr", descr)
            doc_pr.set("title", descr)


def main():
    doc = Document(str(REF))
    clear_body(doc)
    set_core_properties(doc)

    # Cover page uses the reference document's A4 section, margins and footer system.
    add_text(doc, "LAPORAN APLIKASI", align=WD_ALIGN_PARAGRAPH.CENTER, before=44, after=2)
    cover = doc.paragraphs[-1]
    for run in cover.runs:
        set_run(run, size=15, bold=True, color=(35, 78, 155))
    add_text(doc, "PROTOTIPE HIGERTECH LIVE MONITORING MAP", align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    for run in doc.paragraphs[-1].runs:
        set_run(run, size=19, bold=True, color=(25, 38, 68))
    add_text(doc, "WebGIS Monitoring Telemetri Indonesia", align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    for run in doc.paragraphs[-1].runs:
        set_run(run, size=11, italic=True, color=(90, 105, 130))

    logo = ROOT / "public/images/brand/higertech-logo.png"
    if logo.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(logo), width=Inches(2.65))
        style_paragraph(p, after=20)

    cover_rows = [
        ("Nama proyek", "Higertech Live Monitoring Map"),
        ("Jenis dokumen", "Laporan pengembangan prototype aplikasi"),
        ("Teknologi", "Laravel, Blade, Vite, Tailwind CSS, Leaflet.js"),
        ("Sumber basemap", "OpenStreetMap dan penyedia tile non-Google"),
        ("Sumber data", "Dummy lokal berbasis SQLite untuk tahap MVP"),
        ("Status", "Prototype / MVP"),
        ("Tanggal", "21 Agustus 2026"),
    ]
    add_table(doc, ["Keterangan", "Nilai"], cover_rows, [1.55, 5.02], font_size=9)
    add_text(doc, "Dokumen ini disusun untuk mendokumentasikan analisis, perancangan, implementasi, pengujian, dan penggunaan prototype map monitoring.", align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=0)
    for run in doc.paragraphs[-1].runs:
        set_run(run, size=9, color=(75, 85, 105))
    doc.add_page_break()

    add_heading(doc, "KATA PENGANTAR", 1)
    add_text(doc, "Laporan ini mendokumentasikan pengembangan prototype Higertech Live Monitoring Map sebagai halaman WebGIS untuk memantau jaringan stasiun telemetri di Indonesia. Aplikasi dibuat dengan pendekatan MVP sehingga fokusnya berada pada visualisasi peta, pencarian, filter, ringkasan status, dan detail telemetry station.")
    add_text(doc, "Struktur laporan menggunakan pola laporan aplikasi, sedangkan tata letak dokumen, ukuran halaman, gaya heading, caption gambar, dan sistem footer mengikuti template Manual Book User SIH3 Maluku yang menjadi acuan. Data yang digunakan pada prototype adalah data dummy lokal dan bukan salinan data operasional Higertech.")
    add_text(doc, "Semoga dokumen ini membantu proses review desain dan menjadi dasar untuk pengembangan fase berikutnya.")
    doc.add_page_break()

    add_heading(doc, "DAFTAR ISI", 1)
    toc_items = [
        "BAB I PENDAHULUAN",
        "  1.1 Latar Belakang",
        "  1.2 Identifikasi Masalah",
        "  1.3 Tujuan Pengembangan",
        "  1.4 Batasan Pengembangan",
        "  1.5 Sasaran Pengguna",
        "BAB II ANALISIS KEBUTUHAN APLIKASI",
        "  2.1 Deskripsi Umum Aplikasi",
        "  2.2 Kebutuhan Fungsional",
        "  2.3 Kebutuhan Nonfungsional",
        "  2.4 Data dan Klasifikasi Station",
        "  2.5 Skenario Pengguna",
        "BAB III PERANCANGAN SISTEM",
        "  3.1 Arsitektur Teknologi",
        "  3.2 Perancangan Data Station",
        "  3.3 Perancangan Endpoint API",
        "  3.4 Perancangan UI/UX",
        "  3.5 Alur Interaksi",
        "BAB IV IMPLEMENTASI APLIKASI",
        "  4.1 Struktur Project",
        "  4.2 Halaman Live Monitoring Map",
        "  4.3 Search dan Autocomplete",
        "  4.4 Filter Station dan Instansi",
        "  4.5 Marker, Cluster, dan Popup Telemetry",
        "  4.6 Basemap dan Kontrol Peta",
        "  4.7 Loading, Empty, dan Error State",
        "  4.8 API dan Perlindungan Data",
        "BAB V PENGUJIAN",
        "BAB VI PANDUAN PENGGUNAAN",
        "BAB VII KESIMPULAN DAN PENGEMBANGAN LANJUTAN",
        "LAMPIRAN A DAN LAMPIRAN B",
    ]
    for item in toc_items:
        p = add_text(doc, item, after=1)
        for run in p.runs:
            set_run(run, size=9.5, bold=item.startswith("BAB") or item.startswith("LAMPIRAN"))
    doc.add_page_break()

    # BAB I
    add_heading(doc, "BAB I PENDAHULUAN", 1)
    add_heading(doc, "1.1 Latar Belakang", 2)
    add_text(doc, "Pemantauan station hidrologi dan meteorologi membutuhkan tampilan yang dapat membantu pengguna memahami lokasi, jenis perangkat, serta status jaringan secara cepat. Tampilan peta yang hanya berisi marker tanpa konteks akan menyulitkan pengguna ketika jumlah station bertambah, sedangkan informasi yang terlalu padat dapat mengurangi keterbacaan.")
    add_text(doc, "Prototype Higertech Live Monitoring Map dikembangkan untuk menyediakan satu halaman monitoring yang berfokus pada peta Indonesia, ringkasan status jaringan, pencarian station, filter, daftar station, dan popup telemetry. Desainnya mengambil tema dashboard WebGIS modern, tetapi tetap mempertahankan identitas navbar Higertech dan menggunakan komponen visual yang dibuat khusus untuk prototype.")
    add_heading(doc, "1.2 Identifikasi Masalah", 2)
    add_label_para(doc, "Masalah utama. ", "Pengguna membutuhkan cara yang lebih cepat untuk menemukan station tertentu dan membedakan jenis station di peta.")
    add_label_para(doc, "Masalah informasi. ", "Status online/offline, instansi, lokasi, dan pembacaan terbaru perlu disajikan dalam hierarki informasi yang jelas.")
    add_label_para(doc, "Masalah integrasi. ", "Prototype perlu berjalan tanpa ketergantungan pada Google Maps API, API key milik pihak lain, atau endpoint operasional Higertech.")
    add_label_para(doc, "Masalah ketahanan. ", "Data null, station tanpa koordinat, API error, dan hasil filter kosong tidak boleh membuat seluruh halaman gagal digunakan.")
    add_heading(doc, "1.3 Tujuan Pengembangan", 2)
    add_text(doc, "Tujuan pengembangan adalah menghasilkan prototype interaktif pada route /map yang dapat digunakan sebagai gambaran awal fitur Live Map. Prototype ini menyediakan peta Leaflet, basemap OpenStreetMap, marker station dengan warna dan ikon per tipe, clustering, summary card, sidebar pencarian, filter, daftar station, popup telemetry, serta endpoint JSON yang sederhana.")
    add_heading(doc, "1.4 Batasan Pengembangan", 2)
    add_text(doc, "Ruang lingkup dibatasi pada halaman map dan API station. Data bersifat dummy lokal dengan SQLite; belum ada login, admin panel, CMS, CRUD station, halaman detail terpisah, histori grafik, websocket, MQTT, maupun integrasi realtime operasional. Koordinat yang kosong tidak dibuat marker, dan field telemetry yang tidak tersedia tidak ditampilkan sebagai baris kosong.")
    add_heading(doc, "1.5 Sasaran Pengguna", 2)
    add_text(doc, "Sasaran pengguna utama adalah operator monitoring, pengelola balai atau instansi, reviewer produk, dan stakeholder yang membutuhkan gambaran sebaran station. Pada tahap prototype, pengguna tidak perlu melakukan autentikasi untuk mencoba alur peta.")

    # BAB II
    add_heading(doc, "BAB II ANALISIS KEBUTUHAN APLIKASI", 1)
    add_heading(doc, "2.1 Deskripsi Umum Aplikasi", 2)
    add_text(doc, "Aplikasi merupakan halaman WebGIS monitoring telemetri Indonesia. Peta menjadi area utama, sementara panel mengambang di sisi kiri memisahkan area Temukan Pos dan Daftar Pos. Summary ringkas diletakkan di atas map agar status jaringan tetap terlihat tanpa menutup area peta secara berlebihan.")
    add_heading(doc, "2.2 Kebutuhan Fungsional", 2)
    functional_rows = [
        ("F-01", "Menampilkan peta", "Peta Indonesia dengan marker station yang memiliki koordinat.", "Tersedia"),
        ("F-02", "Mencari station", "Input nama station memberi rekomendasi autocomplete dan dapat membuka popup.", "Tersedia"),
        ("F-03", "Memfilter station", "Filter tipe, status, dan instansi/balai dapat diterapkan dan di-reset.", "Tersedia"),
        ("F-04", "Melihat daftar station", "Card station menampilkan nama, tipe, balai, lokasi, update, dan status.", "Tersedia"),
        ("F-05", "Melihat telemetry", "Klik marker/card membuka popup dengan telemetry sesuai tipe station.", "Tersedia"),
        ("F-06", "Mengganti basemap", "Pengguna dapat memilih beberapa basemap non-Google.", "Tersedia"),
        ("F-07", "Melihat ringkasan", "Total, jenis utama, instansi, online, dan offline ditampilkan sebagai summary.", "Tersedia"),
        ("F-08", "Mengakses API", "Data station dan ringkasan tersedia melalui endpoint JSON internal aplikasi.", "Tersedia"),
    ]
    add_table(doc, ["ID", "Fitur", "Kriteria", "Status"], functional_rows, [0.55, 1.35, 3.55, 1.12], font_size=8.3)
    add_heading(doc, "2.3 Kebutuhan Nonfungsional", 2)
    add_label_para(doc, "Usability. ", "Peta menjadi fokus utama, kontrol penting mudah ditemukan, panel dapat ditutup, dan label menggunakan bahasa yang familiar bagi operator.")
    add_label_para(doc, "Responsivitas. ", "Layout desktop memakai panel mengambang; pada viewport kecil panel berubah menjadi drawer dengan ukuran yang lebih terkontrol.")
    add_label_para(doc, "Ketahanan data. ", "Null latest_reading, wilayah kosong, koordinat null, API error, dan data kosong ditangani dengan fallback aman.")
    add_label_para(doc, "Kinerja. ", "Marker dikelompokkan ketika berdekatan, data publik dibatasi pada kolom yang diperlukan, dan endpoint memakai throttle.")
    add_label_para(doc, "Kemandirian sumber. ", "Prototype memakai data dummy lokal dan tidak memanggil endpoint GetStationAll milik website referensi.")
    add_heading(doc, "2.4 Data dan Klasifikasi Station", 2)
    add_text(doc, "Struktur data station mengikuti pola umum API telemetri, tetapi isinya dibuat sendiri untuk prototype. Tipe utama yang ditonjolkan pada UI adalah ARR (curah hujan), AWLR (duga air), AWS (klimatologi), dan AWLR_ARR (duga air serta curah hujan). Tipe lain tetap tersedia dan ditampilkan pada kelompok Lainnya.")
    type_rows = [
        ("ARR", "Pos Curah Hujan", "rainfall, rainfall_last_hour, intensity"),
        ("AWLR", "Pos Duga Air", "water_level, warning_status"),
        ("AWS", "Pos Klimatologi", "temperature, humidity, pressure, wind"),
        ("AWLR_ARR", "Duga Air + Curah Hujan", "water_level, rainfall, warning_status"),
        ("AGWLR", "Air Tanah", "telemetry tersedia jika ada"),
        ("FM", "Flow Meter", "flow_rate, flow_total, flow_month"),
        ("EWS, AVWR, WQ, VNOTCH, OW, OSP", "Lainnya", "ditampilkan dengan label tipe yang sesuai"),
    ]
    add_table(doc, ["Tipe", "Kategori UI", "Contoh telemetry"], type_rows, [1.65, 2.0, 2.92], font_size=8.3)
    add_heading(doc, "2.5 Skenario Pengguna", 2)
    add_text(doc, "Pengguna membuka /map, membaca summary, lalu mencari station melalui kolom Cari pos. Saat rekomendasi muncul, pengguna memilih station untuk memindahkan peta dan membuka popup. Jika ingin mempersempit hasil, pengguna membuka Filter station, memilih tipe/status/instansi, menekan Terapkan Filter, kemudian meninjau daftar card dan marker yang tersisa.")

    # BAB III
    add_heading(doc, "BAB III PERANCANGAN SISTEM", 1)
    add_heading(doc, "3.1 Arsitektur Teknologi", 2)
    add_text(doc, "Aplikasi menggunakan Laravel sebagai backend dan Blade sebagai view. Vite memproses asset frontend, Tailwind CSS digunakan untuk token dan utilitas visual, sedangkan JavaScript modular mengelola state map, basemap, filter, marker, popup, dan kontrol. Leaflet menjadi engine peta; OpenStreetMap serta tile non-Google menjadi pilihan basemap.")
    add_rich(doc, [("Browser", {"bold": True}), (" menampilkan Blade, asset Vite, dan map Leaflet. ", {}), ("Laravel", {"bold": True}), (" menyediakan route halaman, controller API, resource JSON, service query, dan validasi. ", {}), ("SQLite", {"bold": True}), (" menyimpan dummy station untuk lingkungan prototype.", {})])
    add_heading(doc, "3.2 Perancangan Data Station", 2)
    add_text(doc, "Model Station memiliki field identitas station, koordinat, organisasi, wilayah, perangkat, status, waktu pembacaan, dan latest_reading JSON. Kolom publik yang dikirim ke frontend dibatasi pada kebutuhan map. Device ID dimasking menjadi format DEVICE-***-suffix agar contoh data tidak menampilkan identitas perangkat secara penuh.")
    station_rows = [
        ("Identitas", "id, name, slug, station_type", "Nama dan klasifikasi station"),
        ("Lokasi", "latitude, longitude, province_name, regency_name", "Marker, pencarian lokasi, dan popup"),
        ("Instansi", "balai_name, organization_code", "Filter instansi dan summary organisasi"),
        ("Perangkat", "device_id, device_status, timezone", "Status serta identitas perangkat yang dimasking"),
        ("Telemetry", "reading_at, latest_reading", "Popup data terbaru sesuai tipe station"),
    ]
    add_table(doc, ["Kelompok", "Field", "Kegunaan"], station_rows, [1.15, 2.6, 2.82], font_size=8.3)
    add_heading(doc, "3.3 Perancangan Endpoint API", 2)
    api_rows = [
        ("GET", "/api/stations", "data station berkoordinat", "search, type, status, organization"),
        ("GET", "/api/stations/summary", "total, online, offline, organizations, types", "tanpa query wajib"),
        ("GET", "/map", "halaman Blade map", "route halaman"),
    ]
    add_table(doc, ["Metode", "Endpoint", "Respons", "Parameter"], api_rows, [0.65, 1.75, 2.4, 1.77], font_size=8.3)
    add_text(doc, "Endpoint station menggunakan resource JSON dengan meta count dan organizations. Query station tanpa latitude atau longitude tidak dikirim ke marker layer sehingga tidak mengganggu proses rendering peta.")
    add_heading(doc, "3.4 Perancangan UI/UX", 2)
    add_text(doc, "Navbar mempertahankan pola website utama: topbar berisi sosial dan kontak, branding Higertech, menu Home/Product/Projects/Articles/Download/Peta, logo InaProc, dan toggle bahasa. Area map menggunakan panel putih dengan radius dan shadow tipis agar terasa seperti command center modern tanpa menutupi peta secara berlebihan.")
    add_label_para(doc, "Hierarki informasi. ", "Summary berada di atas; Temukan Pos berada di kiri atas; Daftar Pos berada tepat di bawahnya; legend marker berada di sisi bawah map.")
    add_label_para(doc, "Konsistensi kategori. ", "Setiap tipe station memakai warna, icon, badge, dan label yang sama pada marker, card, legend, dan popup.")
    add_label_para(doc, "Kontrol pengguna. ", "Tombol tutup, reset, fullscreen, zoom, fit bounds, basemap, serta Terapkan Filter diletakkan dekat konteks yang dikendalikan.")
    add_heading(doc, "3.5 Alur Interaksi", 2)
    add_text(doc, "Alur utama dimulai dari fetch summary dan station secara paralel. State frontend kemudian membentuk marker, cluster, daftar card, dan opsi instansi. Perubahan search/filter memanggil ulang endpoint dengan query yang telah diringkas. Klik card atau suggestion memanggil flyTo dan membuka popup. Jika request gagal, state error menampilkan retry tanpa membuat seluruh halaman crash.")

    # BAB IV
    add_heading(doc, "BAB IV IMPLEMENTASI APLIKASI", 1)
    add_heading(doc, "4.1 Struktur Project", 2)
    add_text(doc, "Project dibagi sesuai tanggung jawab Laravel dan frontend. Model, migration, factory, seeder, service, resource, dan controller berada di sisi backend. View Blade dipecah menjadi header, sidebar, summary, dan halaman utama. JavaScript map dipisah menjadi modul state, basemap, controls, constants, formatters, dan popup. Pendekatan ini menjaga file tetap mudah ditinjau tanpa membuat abstraksi berlebihan.")
    add_heading(doc, "4.2 Halaman Live Monitoring Map", 2)
    add_text(doc, "Halaman /map menampilkan navbar, summary status, map Indonesia, panel Temukan Pos, panel Daftar Pos, legend marker, dan tombol Basemap. Panel daftar dibuat cukup lebar untuk membaca nama station, jenis, balai, lokasi, waktu update, serta status online/offline.")
    add_image(doc, WORK / "app-map-default.png", "Gambar 4.1 Halaman utama Live Monitoring Map dengan summary, panel pencarian, daftar station, marker cluster, legend, dan basemap.")
    add_heading(doc, "4.3 Search dan Autocomplete", 2)
    add_text(doc, "Kolom pencarian menggunakan nama station sebagai kata kunci. Ketika pengguna mengetik, frontend membentuk rekomendasi dari data hasil query dan membatasi jumlah suggestion agar panel tetap ringkas. Memilih suggestion akan menutup daftar rekomendasi, melakukan flyTo, serta membuka popup station.")
    add_image(doc, WORK / "app-map-autocomplete.png", "Gambar 4.2 Rekomendasi autocomplete saat pengguna mencari station berdasarkan nama.")
    add_heading(doc, "4.4 Filter Station dan Instansi", 2)
    add_text(doc, "Filter dikelompokkan menjadi tipe station, status perangkat, dan instansi/balai. Pada keadaan terbuka, panel filter mengambil ruang yang cukup agar pilihan dapat dibaca. Hasil station disembunyikan sementara supaya pengguna fokus memilih filter, lalu ditampilkan kembali setelah tombol Terapkan Filter ditekan. Tombol Reset Filter mengembalikan kondisi awal.")
    add_image(doc, WORK / "app-map-filter.png", "Gambar 4.3 Panel filter station yang terbuka dengan pilihan tipe, status, instansi, dan tombol penerapan filter.")
    add_heading(doc, "4.5 Marker, Cluster, dan Popup Telemetry", 2)
    add_text(doc, "Marker dibuat dari latitude dan longitude yang valid. Warna dan icon membedakan seluruh dua belas tipe station yang didukung. Ketika marker terlalu berdekatan, Leaflet MarkerCluster mengelompokkan marker untuk menjaga peta tetap terbaca. Klik marker atau card membuka popup berbentuk card yang menampilkan nama, badge tipe, status, instansi, lokasi, device ID yang dimasking, update terakhir, dan telemetry yang tersedia.")
    add_image(doc, WORK / "app-map-default.png", "Gambar 4.4 Marker station dan cluster pada peta; warna marker membedakan tipe station dan cluster menjaga peta tetap terbaca pada zoom rendah.")
    add_heading(doc, "4.6 Basemap dan Kontrol Peta", 2)
    add_text(doc, "Basemap switcher menyediakan enam pilihan: OpenStreetMap, Humanitarian, Topografi, Light, Dark, dan Sentinel-2 Satellite. Semuanya dipilih sebagai alternatif non-Google. Kontrol peta mencakup zoom, reset view, fit bounds, fullscreen jika tersedia, serta panel basemap mengambang di sisi map.")
    add_image(doc, WORK / "app-map-default.png", "Gambar 4.5 Basemap dan kontrol peta pada area map, termasuk zoom, reset view, fullscreen, fit bounds, dan tombol Basemap.")
    add_heading(doc, "4.7 Loading, Empty, dan Error State", 2)
    add_text(doc, "Saat request berjalan, panel menampilkan loading. Jika hasil filter kosong, pengguna menerima empty state yang menjelaskan bahwa tidak ada station yang cocok. Jika API error, halaman menampilkan pesan singkat dan tombol retry. Station dengan data wilayah atau latest_reading null tetap dapat dirender karena formatter memakai fallback aman dan hanya menampilkan field telemetry yang tersedia.")
    add_image(doc, WORK / "app-map-filter.png", "Gambar 4.6 Panel filter dalam keadaan aktif; area hasil dapat dikosongkan atau dimuat ulang tanpa membuat map utama berhenti digunakan.")
    add_heading(doc, "4.8 API dan Perlindungan Data", 2)
    add_text(doc, "Frontend mengakses endpoint internal aplikasi, bukan endpoint operasional website referensi. Controller mengembalikan resource yang sudah dipilih, service membatasi query publik, dan device ID dimasking. Throttle station-api digunakan untuk mencegah request berlebihan. Perlu dipahami bahwa endpoint yang dipanggil browser tetap dapat terlihat melalui Network tab; perlindungan yang realistis adalah membatasi data publik dan menambahkan autentikasi ketika aplikasi masuk fase produksi, bukan menyembunyikan API secara absolut.")

    # BAB V
    add_heading(doc, "BAB V PENGUJIAN", 1)
    add_heading(doc, "5.1 Strategi Pengujian", 2)
    add_text(doc, "Pengujian dilakukan pada tiga lapisan: pengujian backend Laravel/PHP, pengujian JavaScript state, dan pengujian build frontend. Selain itu dilakukan pemeriksaan browser pada route /map untuk memastikan peta, panel, suggestion, dan filter dapat digunakan secara visual.")
    add_heading(doc, "5.2 Pengujian Fungsional", 2)
    test_rows = [
        ("T-01", "Buka /map", "Navbar, map, summary, panel, dan marker tampil.", "Lulus"),
        ("T-02", "Cari Air Tanah", "Suggestion station muncul dan dapat dipilih.", "Lulus"),
        ("T-03", "Buka filter", "Kontrol filter terbuka, daftar hasil ditutup sementara.", "Lulus"),
        ("T-04", "Terapkan filter", "Marker dan card mengikuti tipe/status/instansi.", "Lulus"),
        ("T-05", "Klik card/marker", "Peta flyTo dan popup telemetry terbuka.", "Lulus"),
        ("T-06", "Reset view/basemap", "Peta kembali ke view dan layer yang dipilih.", "Lulus"),
        ("T-07", "Null/error/empty", "Fallback aman, empty state, dan retry tersedia.", "Lulus"),
    ]
    add_table(doc, ["ID", "Skenario", "Hasil yang diharapkan", "Hasil"], test_rows, [0.55, 1.6, 3.3, 1.12], font_size=8.3)
    add_heading(doc, "5.3 Pengujian API", 2)
    add_text(doc, "Endpoint GET /api/stations diuji untuk respons data station berkoordinat serta parameter search, type, status, dan organization. Endpoint GET /api/stations/summary diuji untuk total station, online/offline, jumlah organisasi, dan distribusi tipe. Struktur resource juga diperiksa agar latest_reading nullable dan device ID tidak dikirim secara penuh.")
    add_heading(doc, "5.4 Pengujian Frontend dan Build", 2)
    add_label_para(doc, "PHP/Laravel. ", "17 test lulus dengan 207 assertions pada pemeriksaan terakhir.")
    add_label_para(doc, "JavaScript. ", "21 test lulus, termasuk state filter, suggestion, dan helper map.")
    add_label_para(doc, "Code style. ", "Laravel Pint selesai tanpa pelanggaran pada source yang diperiksa.")
    add_label_para(doc, "Build. ", "Vite build berhasil menghasilkan asset production.")
    add_label_para(doc, "Browser smoke check. ", "Route /map berhasil menampilkan peta, marker, autocomplete, dan panel filter pada browser lokal.")
    add_heading(doc, "5.5 Ringkasan Hasil", 2)
    add_text(doc, "Berdasarkan pengujian tersebut, prototype memenuhi kebutuhan MVP yang ditetapkan untuk map interaktif. Pengujian produksi dengan data operasional dan pengujian beban belum termasuk dalam scope dokumen ini dan perlu dilakukan pada fase integrasi berikutnya.")

    # BAB VI
    add_heading(doc, "BAB VI PANDUAN PENGGUNAAN", 1)
    add_heading(doc, "6.1 Menjalankan Aplikasi", 2)
    add_text(doc, "Pastikan PHP, Composer, Node.js, dan npm tersedia. Dari root project jalankan perintah berikut:")
    add_rich(doc, [("composer install", {"bold": True}), (" untuk dependency PHP; ", {}), ("npm install", {"bold": True}), (" untuk dependency frontend; ", {}), ("cp .env.example .env", {"bold": True}), (" lalu ", {}), ("php artisan key:generate", {"bold": True}), ("; setelah itu ", {}), ("php artisan migrate:fresh --seed", {"bold": True}), ("; dan ", {}), ("npm run build", {"bold": True}), (". Jalankan server dengan ", {}), ("php artisan serve", {"bold": True}), (".", {})])
    add_heading(doc, "6.2 Membuka Halaman Map", 2)
    add_text(doc, "Buka http://127.0.0.1:8000/map. Tunggu sampai summary dan station selesai dimuat. Jika jaringan tile belum tersedia, panel error akan memberi informasi dan peta dapat dicoba kembali setelah koneksi tersedia.")
    add_heading(doc, "6.3 Mencari Station", 2)
    add_text(doc, "Klik kolom Cari pos, ketik sebagian nama station, lalu pilih salah satu rekomendasi. Peta akan mendekat ke lokasi station dan popup telemetry akan terbuka. Tekan Escape untuk menutup rekomendasi tanpa menutup panel pencarian.")
    add_heading(doc, "6.4 Menerapkan Filter", 2)
    add_text(doc, "Buka Filter station, pilih tipe station, status, dan/atau instansi. Tekan Terapkan Filter untuk memuat hasil. Jika ingin kembali ke seluruh station, tekan Reset Filter. Saat filter terbuka, daftar hasil disembunyikan sementara supaya fokus interaksi tetap jelas.")
    add_heading(doc, "6.5 Memilih Card dan Popup", 2)
    add_text(doc, "Klik card pada Daftar Pos untuk memindahkan peta ke station. Klik marker atau cluster untuk melihat detail. Popup hanya menampilkan telemetry yang tersedia; nilai null tidak dibuat menjadi baris kosong.")
    add_heading(doc, "6.6 Mengubah Basemap dan Kontrol Peta", 2)
    add_text(doc, "Gunakan tombol Basemap untuk mengganti gaya peta. Gunakan zoom, reset, fit bounds, dan fullscreen sesuai kebutuhan. Pada layar kecil, panel pencarian dan daftar dapat ditutup dengan tombol X sehingga area peta kembali luas.")
    add_heading(doc, "6.7 Interpretasi Status", 2)
    add_text(doc, "Badge hijau menunjukkan station online dan badge merah menunjukkan offline. Summary di bagian atas merangkum kondisi seluruh data yang sedang tersedia. Status ini berasal dari dummy seed pada prototype dan belum merepresentasikan kondisi operasional nyata.")

    # BAB VII
    add_heading(doc, "BAB VII KESIMPULAN DAN PENGEMBANGAN LANJUTAN", 1)
    add_heading(doc, "7.1 Kesimpulan", 2)
    add_text(doc, "Prototype Higertech Live Monitoring Map telah membentuk fondasi halaman map interaktif yang berfokus pada keterbacaan dan alur kerja operator. Pengguna dapat melihat sebaran station Indonesia, membaca ringkasan status, mencari station dengan autocomplete, menerapkan filter, membuka popup telemetry, dan mengganti basemap non-Google. Backend dan frontend dipisah secara modular, sedangkan fallback null/error/empty menjaga halaman tetap stabil.")
    add_heading(doc, "7.2 Rekomendasi Fase Berikutnya", 2)
    add_text(doc, "Setelah desain MVP disetujui, pengembangan berikutnya dapat mencakup integrasi sumber data resmi melalui backend, autentikasi dan otorisasi, admin panel station, histori telemetry dan grafik, notifikasi status, WebSocket/MQTT, logging, cache, serta pengujian beban. Integrasi data operasional sebaiknya dilakukan setelah kontrak field, keamanan, dan hak akses ditetapkan.")

    add_heading(doc, "LAMPIRAN A RINGKASAN API", 1)
    add_text(doc, "Contoh bentuk respons /api/stations adalah objek JSON dengan data station dan meta count/organizations. Field utama yang digunakan frontend meliputi id, name, station_type, latitude, longitude, balai_name, province_name, regency_name, device_id, device_status, reading_at, dan latest_reading. Endpoint summary mengembalikan total, online, offline, organizations, dan types.")
    add_heading(doc, "LAMPIRAN B GLOSARIUM TIPE STATION", 1)
    add_text(doc, "ARR adalah Automatic Rainfall Recorder atau pos curah hujan; AWLR adalah Automatic Water Level Recorder atau pos duga air; AWS adalah Automatic Weather Station atau pos klimatologi; AWLR_ARR menggabungkan duga air dan curah hujan; AGWLR digunakan untuk air tanah; FM untuk flow meter; EWS untuk early warning system; AVWR untuk automatic valve; WQ untuk kualitas air; VNOTCH untuk v-notch; OW untuk observation well; dan OSP untuk outlet structure pump.")
    add_text(doc, "Dokumen ini merupakan laporan prototype. Angka station, koordinat, status, telemetry, dan screenshot yang ditampilkan berasal dari seed dummy lokal untuk kebutuhan demonstrasi.", before=10)

    set_alt_text_for_drawings(doc)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
